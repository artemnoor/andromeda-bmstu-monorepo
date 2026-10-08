from __future__ import annotations

import re
import zipfile
from collections.abc import Iterable
from typing import Any
from urllib.parse import urldefrag, urljoin, urlparse

from bs4 import BeautifulSoup
from pypdf.errors import PdfReadError

from ..common import clean_text

OFFICIAL_RSOSH_2025_26_URL = "https://rsr-olymp.ru/archive"
RSOSH_HOST = "rsr-olymp.ru"
PROFILE_LINK_HINTS = (
    "профил", "направлен", "предмет", "положен", "регламент", "документ", "список", "программа",
)
DOCUMENT_SUFFIXES = {".pdf", ".xls", ".xlsx", ".doc", ".docx", ".odt", ".ods"}


def normalize_event_name(value: str) -> str:
    return re.sub(r"[^0-9a-zа-яё]+", "", clean_text(value).casefold())


def parse_rsosh_directory(body: bytes, source_url: str = OFFICIAL_RSOSH_2025_26_URL) -> list[dict[str, Any]]:
    """Read organizer URLs linked from the official RSOSH 2025/26 index."""
    soup = BeautifulSoup(body, "html.parser")
    rows: list[dict[str, Any]] = []
    for table_no, table in enumerate(soup.find_all("table"), start=1):
        table_rows = table.find_all("tr")
        first_row = [clean_text(cell.get_text(" ", strip=True)) for cell in table_rows[0].find_all(["td", "th"], recursive=False)] if table_rows else []
        if not any("название" in cell.casefold() for cell in first_row) or not any("профиль" in cell.casefold() for cell in first_row):
            continue
        current_title = ""
        current_url: str | None = None
        for row_no, row in enumerate(table_rows[1:], start=2):
            cells = [clean_text(cell.get_text(" ", strip=True)) for cell in row.find_all(["td", "th"], recursive=False)]
            if not cells or not any(cells):
                continue
            title_cell = cells[1] if len(cells) >= 5 and cells[0].isdigit() else None
            if title_cell:
                current_title = title_cell
                anchor = row.find("a", href=True)
                href = clean_text(anchor.get("href")) if anchor else ""
                candidate_url = urljoin(source_url, href) if href else ""
                parsed = urlparse(candidate_url)
                current_url = candidate_url if parsed.scheme == "https" and parsed.hostname and parsed.hostname.casefold().rstrip(".") != RSOSH_HOST else None
            elif len(cells) >= 5 and cells[1]:
                current_title = cells[1]
                anchor = row.find("a", href=True)
                href = clean_text(anchor.get("href")) if anchor else ""
                candidate_url = urljoin(source_url, href) if href else ""
                parsed = urlparse(candidate_url)
                current_url = candidate_url if parsed.scheme == "https" and parsed.hostname and parsed.hostname.casefold().rstrip(".") != RSOSH_HOST else None
            if not current_title:
                continue
            # RSOSH uses rowspans: continuation profile rows contain only profile, subject, and level.
            profile = cells[2] if len(cells) >= 5 else cells[0] if len(cells) >= 3 else None
            subject = cells[3] if len(cells) >= 5 else cells[1] if len(cells) >= 3 else None
            level = cells[4] if len(cells) >= 5 else cells[2] if len(cells) >= 3 else None
            if not profile or profile.casefold() == "профиль":
                continue
            rows.append({
                "external_key": f"olympiad_organizer_directory:rsosh:2025-26:t{table_no}:r{row_no}",
                "event_name": current_title,
                "normalized_event_name": normalize_event_name(current_title),
                "organizer_official_url": current_url,
                "profile_name_in_rsosh_directory": profile,
                "subject_in_rsosh_directory": subject,
                "level_in_rsosh_directory": level,
                "rsosh_row_cells": cells,
                "source_url": source_url,
                "rsosh_source_url": source_url,
            })
    return rows


def official_related_links(body: bytes, source_url: str, *, profile_names: Iterable[str] = (), limit: int = 6) -> list[dict[str, str]]:
    """Select explicit profile/rules/document links from the linked official site."""
    soup = BeautifulSoup(body, "html.parser")
    base = urlparse(source_url)
    sibling_hosts = {base.hostname.casefold()} if base.hostname else set()
    if base.hostname and base.hostname.casefold().startswith("www."):
        sibling_hosts.add(base.hostname.casefold()[4:])
    elif base.hostname:
        sibling_hosts.add("www." + base.hostname.casefold())
    candidates: dict[str, tuple[int, str]] = {}
    normalized_profiles = [normalize_event_name(value) for value in profile_names if normalize_event_name(value)]
    for anchor in soup.find_all("a", href=True):
        href = clean_text(anchor.get("href"))
        if not href or href.startswith(("#", "mailto:", "tel:", "javascript:")):
            continue
        url, _fragment = urldefrag(urljoin(source_url, href))
        parsed = urlparse(url)
        if parsed.scheme != "https" or not parsed.hostname or parsed.hostname.casefold() not in sibling_hosts:
            continue
        if url.rstrip("/") == source_url.rstrip("/"):
            continue
        label = clean_text(anchor.get_text(" ", strip=True))
        target_text = f"{label} {parsed.path}".casefold()
        normalized_target = normalize_event_name(target_text)
        suffix = next((suffix for suffix in DOCUMENT_SUFFIXES if parsed.path.casefold().endswith(suffix)), "")
        hint = any(value in target_text for value in PROFILE_LINK_HINTS)
        profile_match = any(value in normalized_target for value in normalized_profiles)
        if not hint and not profile_match:
            continue
        score = (5 if profile_match else 0) + (3 if suffix else 0) + (1 if hint else 0)
        candidates[url] = (max(score, candidates.get(url, (0, ""))[0]), label)
    ordered = sorted(candidates.items(), key=lambda item: (-item[1][0], item[0]))[:limit]
    return [{"url": url, "link_text": label} for url, (_score, label) in ordered]


def extract_official_page_text(body: bytes, url: str) -> str:
    """Extract searchable text from common organizer pages and linked documents."""
    suffix = urlparse(url).path.casefold().rsplit(".", 1)[-1] if "." in urlparse(url).path else ""
    if suffix == "pdf":
        try:
            import io

            from pypdf import PdfReader

            return " ".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(body)).pages)
        except (ImportError, OSError, ValueError, RuntimeError, TypeError, PdfReadError):
            return ""
    if suffix == "xlsx":
        try:
            import io

            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            return " ".join(str(value) for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
        except (ImportError, OSError, ValueError, TypeError, KeyError, IndexError, zipfile.BadZipFile):
            return ""
        except InvalidFileException:
            return ""
    if suffix == "docx":
        try:
            import io
            from xml.etree import ElementTree

            with zipfile.ZipFile(io.BytesIO(body)) as archive:
                xml = archive.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            return " ".join(text for text in root.itertext() if text)
        except (ImportError, OSError, ValueError, TypeError, KeyError, zipfile.BadZipFile):
            return ""
        except ElementTree.ParseError:
            return ""
    return clean_text(BeautifulSoup(body, "html.parser").get_text(" ", strip=True))


__all__ = [
    "OFFICIAL_RSOSH_2025_26_URL",
    "extract_official_page_text",
    "normalize_event_name",
    "official_related_links",
    "parse_rsosh_directory",
]
