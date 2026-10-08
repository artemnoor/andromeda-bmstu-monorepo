from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from typing import Any

from ....curriculum_identity import reconcile_curriculum_rows
from ..common import normalize_code


def parse_curriculum_document(
    body: bytes,
    *,
    content_type: str | None,
    final_url: str,
    share_url: str,
    profile: dict[str, Any],
    retrieved_at: str,
) -> dict[str, Any]:
    """Parse one document selected by one exact profile card link."""
    if body.startswith(b"%PDF-") or (content_type or "").lower().find("pdf") >= 0:
        parsed = _parse_pdf(body, final_url, share_url, profile, retrieved_at)
        return _attach_exact_identity(
            parsed,
            body=body,
            final_url=final_url,
            share_url=share_url,
            profile=profile,
        )
    if body.startswith(b"PK\x03\x04") or any(token in (content_type or "").lower() for token in ("spreadsheet", "excel", "officedocument")):
        return _parse_xlsx(body, final_url, share_url, profile, retrieved_at)
    return {
        "format": "unsupported",
        "profile_code": profile["code"],
        "header_profile_code": None,
        "education_year": None,
        "duration": None,
        "items": [],
        "error": "linked study-plan file is neither a recognized PDF nor XLS/XLSX workbook",
    }


def _attach_exact_identity(
    parsed: dict[str, Any],
    *,
    body: bytes,
    final_url: str,
    share_url: str,
    profile: dict[str, Any],
) -> dict[str, Any]:
    """Attach stable importer keys only when the selected plan has exact identity."""
    items = [item for item in parsed.get("items", []) if isinstance(item, dict)]
    plan_key = profile.get("study_plan_external_key")
    parsed["source_sha256"] = sha256(body).hexdigest()
    parsed["source_url"] = share_url
    parsed["download_url"] = final_url
    if not isinstance(plan_key, str) or not plan_key.startswith("study_plan:"):
        parsed["identity_status"] = "source_gap"
        parsed["identity_gap"] = "exact_study_plan_external_key_unresolved"
        return parsed

    source_rows: list[dict[str, Any]] = []
    for position, item in enumerate(items, start=1):
        item["parsed_position"] = position
        item["source_sha256"] = parsed["source_sha256"]
        item["source_url"] = share_url
        item["source_locator"] = {
            "page": item.get("source_page"),
            "printed_row_no": item.get("row_no"),
            "parsed_position": position,
            "semester": item.get("semester"),
            "source_url": share_url,
        }
        source_rows.append({
            **item,
            "printed_row_no": item.get("row_no"),
            "parsed_position": position,
        })
    reconciled = reconcile_curriculum_rows(plan_key, source_rows, [])
    for item, resolution in zip(items, reconciled.rows, strict=True):
        item["external_key"] = resolution.canonical_external_key
        item["curriculum_key"] = plan_key
        item["source_position"] = item["parsed_position"]
        identity = {
            key: resolution.metadata.get(key)
            for key in ("algorithm", "signals", "duplicate_title_semester")
            if resolution.metadata.get(key) is not None
        }
        if resolution.metadata.get("source_identity_key") is not None:
            identity["source_identity_id"] = resolution.metadata["source_identity_key"]
        if resolution.status == "ambiguous":
            identity["status"] = "ambiguous"
            item["suggested_existing_keys"] = list(resolution.suggested_existing_keys)
        item["source_row"] = {"identity": identity}
        item["identity_status"] = resolution.status

    if all(item.get("identity_status") != "ambiguous" for item in items):
        parsed["identity_status"] = "exact"
        parsed.pop("identity_gap", None)
    else:
        parsed["identity_status"] = "requires_review"
        parsed["identity_gap"] = "curriculum_rows_need_exact_source_identity_review"
    return parsed


def _parse_pdf(body: bytes, final_url: str, share_url: str, profile: dict[str, Any], retrieved_at: str) -> dict[str, Any]:
    from andromeda.ingestion.universities.bmstu.parser.curriculum import (
        _study_plan_records,
    )
    from andromeda.ingestion.universities.bmstu.source_models import (
        FetchedResource,
        SourceDefinition,
    )

    resource = FetchedResource(
        requested_url=share_url,
        final_url=final_url,
        status_code=200,
        content_type="application/pdf",
        body=body,
        fetched_at=retrieved_at,
    )
    source = SourceDefinition(id="BMSTU-2026", name="BMSTU profile-specific study plan", url=share_url)
    records = _study_plan_records(
        source,
        resource,
        retrieved_at,
        context={
            "program_profile_code": profile["code"],
            "study_plan_url": share_url,
            "document_url": share_url,
            "download_url": final_url,
        },
    )
    summary = next((item for item in records if item.get("record_type") == "StudyPlanSummary"), None)
    items = [item for item in records if item.get("record_type") == "StudyPlan"]
    header_code = normalize_code(summary.get("program_profile_code")) if summary else None
    return {
        "format": "pdf",
        "profile_code": profile["code"],
        "header_profile_code": header_code,
        "direction_code": normalize_code(summary.get("direction_code")) if summary else None,
        "profile_name_in_plan": summary.get("program_profile") if summary else None,
        "education_year": summary.get("education_year") if summary else None,
        "duration": summary.get("duration") if summary else None,
        "study_form": summary.get("form") if summary else None,
        "faculty": summary.get("faculty") if summary else None,
        "department_in_plan": summary.get("department") if summary else None,
        "semester_count": summary.get("semester_count") if summary else None,
        "semester_weeks": summary.get("semester_weeks") if summary else None,
        "semester_totals": summary.get("semesters") if summary else [],
        "overall_astronomical": summary.get("overall_astronomical") if summary else None,
        "overall_academic": summary.get("overall_academic") if summary else None,
        "items": items,
        "error": None if summary and items else "official PDF did not yield a plan summary and curriculum rows",
    }


def _parse_xlsx(body: bytes, final_url: str, share_url: str, profile: dict[str, Any], retrieved_at: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "format": "xlsx",
            "profile_code": profile["code"],
            "header_profile_code": None,
            "education_year": None,
            "duration": None,
            "items": [],
            "error": "openpyxl is unavailable for the linked official workbook",
        }
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    items: list[dict[str, Any]] = []
    sheet_names: list[str] = []
    for sheet in workbook.worksheets:
        sheet_names.append(sheet.title)
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows, start=1):
            values = [str(value).strip() if value is not None else "" for value in row]
            joined = " ".join(values).casefold()
            if not any(term in joined for term in ("дисциплина", "предмет", "семестр")):
                continue
            items.append({
                "record_type": "WorkbookRow",
                "sheet": sheet.title,
                "source_position": index,
                "cells": list(row),
                "profile_code": profile["code"],
            })
    workbook.close()
    return {
        "format": "xlsx",
        "profile_code": profile["code"],
        "header_profile_code": None,
        "education_year": None,
        "duration": None,
        "sheet_names": sheet_names,
        "items": items,
        "error": None if items else "workbook contains no identifiable curriculum headings",
    }
