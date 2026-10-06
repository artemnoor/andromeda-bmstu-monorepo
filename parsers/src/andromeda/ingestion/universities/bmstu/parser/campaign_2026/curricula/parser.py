from __future__ import annotations

from io import BytesIO
from typing import Any

from ..common import normalize_code


def parse_curriculum_document(
    body: bytes,
    *,
    content_type: str | None,
    final_url: str,
    share_url: str,
    profile: dict[str, Any],
    retrieved_at: str,
) -> dict[str, Any]:
    """Parse one document selected by one exact profile card link."""
    if body.startswith(b"%PDF-") or (content_type or "").lower().find("pdf") >= 0:
        return _parse_pdf(body, final_url, share_url, profile, retrieved_at)
    if body.startswith(b"PK\x03\x04") or any(token in (content_type or "").lower() for token in ("spreadsheet", "excel", "officedocument")):
        return _parse_xlsx(body, final_url, share_url, profile, retrieved_at)
    return {
        "format": "unsupported",
        "profile_code": profile["code"],
        "header_profile_code": None,
        "education_year": None,
        "duration": None,
        "items": [],
        "error": "linked study-plan file is neither a recognized PDF nor XLS/XLSX workbook",
    }


def _parse_pdf(body: bytes, final_url: str, share_url: str, profile: dict[str, Any], retrieved_at: str) -> dict[str, Any]:
    from andromeda.ingestion.universities.bmstu.parser.curriculum import _study_plan_records
    from andromeda.ingestion.universities.bmstu.source_models import FetchedResource, SourceDefinition

    resource = FetchedResource(
        requested_url=share_url,
        final_url=final_url,
        status_code=200,
        content_type="application/pdf",
        body=body,
        fetched_at=retrieved_at,
    )
    source = SourceDefinition(id="BMSTU-2026", name="BMSTU profile-specific study plan", url=share_url)
    records = _study_plan_records(
        source,
        resource,
        retrieved_at,
        context={
            "program_profile_code": profile["code"],
            "study_plan_url": share_url,
            "document_url": share_url,
            "download_url": final_url,
        },
    )
    summary = next((item for item in records if item.get("record_type") == "StudyPlanSummary"), None)
    items = [item for item in records if item.get("record_type") == "StudyPlan"]
    header_code = normalize_code(summary.get("program_profile_code")) if summary else None
    return {
        "format": "pdf",
        "profile_code": profile["code"],
        "header_profile_code": header_code,
        "direction_code": normalize_code(summary.get("direction_code")) if summary else None,
        "profile_name_in_plan": summary.get("program_profile") if summary else None,
        "education_year": summary.get("education_year") if summary else None,
        "duration": summary.get("duration") if summary else None,
        "study_form": summary.get("form") if summary else None,
        "faculty": summary.get("faculty") if summary else None,
        "department_in_plan": summary.get("department") if summary else None,
        "semester_count": summary.get("semester_count") if summary else None,
        "semester_weeks": summary.get("semester_weeks") if summary else None,
        "semester_totals": summary.get("semesters") if summary else [],
        "overall_astronomical": summary.get("overall_astronomical") if summary else None,
        "overall_academic": summary.get("overall_academic") if summary else None,
        "items": items,
        "error": None if summary and items else "official PDF did not yield a plan summary and curriculum rows",
    }


def _parse_xlsx(body: bytes, final_url: str, share_url: str, profile: dict[str, Any], retrieved_at: str) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "format": "xlsx",
            "profile_code": profile["code"],
            "header_profile_code": None,
            "education_year": None,
            "duration": None,
            "items": [],
            "error": "openpyxl is unavailable for the linked official workbook",
        }
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    items: list[dict[str, Any]] = []
    sheet_names: list[str] = []
    for sheet in workbook.worksheets:
        sheet_names.append(sheet.title)
        rows = list(sheet.iter_rows(values_only=True))
        for index, row in enumerate(rows, start=1):
            values = [str(value).strip() if value is not None else "" for value in row]
            joined = " ".join(values).casefold()
            if not any(term in joined for term in ("дисциплина", "предмет", "семестр")):
                continue
            items.append({
                "record_type": "WorkbookRow",
                "sheet": sheet.title,
                "source_position": index,
                "cells": list(row),
                "profile_code": profile["code"],
            })
    workbook.close()
    return {
        "format": "xlsx",
        "profile_code": profile["code"],
        "header_profile_code": None,
        "education_year": None,
        "duration": None,
        "sheet_names": sheet_names,
        "items": items,
        "error": None if items else "workbook contains no identifiable curriculum headings",
    }
